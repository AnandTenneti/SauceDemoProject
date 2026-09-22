# utils/common_utils.py

import json
import re
import csv
import openpyxl
from pathlib import Path


class CommonUtils:

    @staticmethod
    def open_xlsx_file(path):
        login_data = []
        workbook = openpyxl.load_workbook(path)
        sheet = workbook.active

        for row in sheet.iter_rows(min_row=2, values_only=True):
            email, password = row

            login_data.append(
                (
                    str(email or ""),
                    str(password or ""),

                )
            )

        workbook.close()
        return login_data

    @staticmethod
    def open_file(path):
        with open(path, encoding="utf-8") as f:
            return json.load(f)

    @staticmethod
    def open_csv_file(path):
        with open(path, "r", newline="") as user_file:
            csv_userfile_reader = csv.DictReader(user_file)
            return [(user["username"], user["password"])
                    for user in csv_userfile_reader]

    @staticmethod
    def get_extension(filename):
        extension = filename.split(".")[-1]
        return "." + extension

    @staticmethod
    @staticmethod
    def open_filetype(path):

        filetype = Path(path).suffix.lower().lstrip(".")

        match filetype:

            case "json":
                print("Reading JSON")
                with open(path, encoding="utf-8") as f:
                    return json.load(f)

            case "csv":
                print("Reading CSV")
                with open(path, "r", newline="", encoding="utf-8") as user_file:
                    csv_userfile_reader = csv.DictReader(user_file)

                    return [
                        (user["username"], user["password"])
                        for user in csv_userfile_reader
                    ]
            case "xlsx":
                login_data = []
                workbook = openpyxl.load_workbook(path)
                sheet = workbook.active

                for row in sheet.iter_rows(min_row=2, values_only=True):
                    email, password = row

                    login_data.append(
                        (
                            str(email or ""),
                            str(password or ""),

                        )
                    )

                workbook.close()
                return login_data

            case _:
                raise ValueError(f"Unsupported file type: {filetype}")

    @staticmethod
    def extract_value(str):
        return float(re.search(r"[\d.]+", str).group())

    @staticmethod
    def format_product_id(text):
        """
        Convert a product name into the format used by SauceDemo
        add-to-cart button IDs.

        Example:
            'Sauce Labs Backpack' -> 'sauce-labs-backpack'

        Args:
            text (str): Product name.

        Returns:
            str: Formatted product identifier.
        """
        return text.replace(" ", "-").lower()
